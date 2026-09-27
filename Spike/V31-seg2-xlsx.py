# V-31 seg2 -- XLSX EXPORT: are the header cells in the PRODUCED FILE byte-identical to
# the engine's labels?
#
# This is a REAL RUN, not a code reading: it calls the same whitelisted endpoint the
# Desk "Export > Excel" button calls (frappe.desk.query_report._export_query), takes the
# xlsx bytes it hands back, writes them under /workspace/Spike/V31-out/, and READS THE
# HEADER ROW BACK with openpyxl -- an independent reader, not frappe's own code.
#
# Why reading the file back matters: query_report.py:692 is
#     column_data.append(_(column.get("label")))
# so the header cells are `_()`-transformed, and `_()` also .strip()s its lookup key.
# Anything xlsxwriter or the sheet-name sanitizer does on top of that would only be
# visible in the bytes. seg1 established what `_()` returns; this seg establishes what
# actually lands in the file.
#
# CRITERION DISCIPLINE: the same four-cell design as seg1, so "the header looked right"
# cannot be confused with "this label had nothing to collide with":
#   NEG  Chinese statutory labels        -> expect byte-identical
#   POS  labels that ARE live zh keys    -> expect the ACCOUNT column SUBSTITUTED while
#        the '<key> - 2026' period column passes through. Two different outcomes inside
#        ONE file prove the readback can actually see a substitution.
# If POS came back unchanged too, the NEG result would prove nothing and this probe
# would have to report that.
#
# SITE SAFETY
# * One transaction, `frappe.db.rollback()` in `finally`. No commit anywhere.
# * `module` left EMPTY on every template (on_update -> _export_template() would write
#   JSON into the erpnext source tree; rollback cannot undo a filesystem write).
# * Pre-flight on the path actually triggered here (grepped, and asserted in seg0):
#     - _export_query itself: `frappe.enqueue` exists at query_report.py:418 but ONLY
#       under `if export_in_background:`. This probe calls _export_query DIRECTLY with
#       export_in_background absent, so that branch is not on the path. Confirmed by
#       recording the enqueue call list via a patch below -- it must stay empty.
#     - frappe.utils.xlsxutils: 0 db.commit, 0 enqueue, 0 save_file (seg0 census).
#     - provide_binary_file (desk/utils.py:80-86) only sets frappe.response keys --
#       no File doctype row, no disk write. This probe calls _export_query with
#       populate_response=False anyway and gets the bytes returned directly.
#     - report.py:192 prepared-report threading.Timer -> enable_prepared_report ->
#       db.commit() ON A SEPARATE CONNECTION: patched to a no-op first, same as seg1.
# * Files: only /workspace/Spike/V31-out/*.xlsx, written by this probe on purpose and
#   reported. Nothing is written to frappe's private files dir (no save_file on path).
# * Error Log is MyISAM (rollback will not undo it): count asserted before/after, and
#   any row present is REPORTED, never deleted by exclusion.
#
# Chinese appears ONLY as \uXXXX escapes -- this source is ASCII-only.
#
# Run:
#   docker exec -i -e PYTHONUTF8=1 -e PYTHONIOENCODING=utf-8 \
#     -w /workspace/frappe-bench/sites erx001-frappe-1 \
#     /workspace/frappe-bench/env/bin/python /workspace/Spike/V31-seg2-xlsx.py

import glob
import hashlib
import inspect
import json
import os
import traceback

import frappe

SITE = "erx.localhost"
COMPANY = "\u534E\u4E1C\u5F39\u7C27"
OUT_DIR = "/workspace/Spike/V31-out"
OUT = os.path.join(OUT_DIR, "seg2-xlsx.json")
REPORT_NAME = "Custom Financial Statement"

TPL_NEG = "ZZ-PROBE-V31X-NEG-StatutoryLiteral"
SEG_NEG_0 = "\u5E74\u521D\u4F59\u989D"
SEG_NEG_1 = "\u671F\u672B\u4F59\u989D"

TPL_LONG = "ZZ-PROBE-V31X-LONG-RealisticPanes"
SEG_LONG_0 = "\u8D44\u4EA7"
SEG_LONG_1 = "\u8D1F\u503A\u548C\u6240\u6709\u8005\u6743\u76CA"

TPL_POS = "ZZ-PROBE-V31X-POS-TranslatedKey"   # display_names chosen at runtime

result = {"probe": "V-31 seg2: XLSX export header cells, read back from the produced file"}


def rec(k, v):
    result[k] = v
    print("[" + k + "] " + json.dumps(v, ensure_ascii=False, default=str)[:2400], flush=True)


def char_dump(s):
    if s is None:
        return None
    return {
        "raw": s,
        "repr": repr(s),
        "len": len(s),
        "codepoints": ["U+%04X" % ord(c) for c in s],
        "has_leading_ws": s != s.lstrip(),
        "has_trailing_ws": s != s.rstrip(),
    }


def bs_rows(seg0_name, seg1_name):
    """Same as V-29 / V-31 seg1, unchanged, so strings stay comparable."""
    asset_f = json.dumps(["root_type", "=", "Asset"])
    liab_f = json.dumps(["root_type", "=", "Liability"])
    eq_f = json.dumps(["root_type", "=", "Equity"])
    return [
        {"data_source": "Column Break", "display_name": seg0_name},
        {"data_source": "Account Data", "display_name": "\u8D44\u4EA7\u5408\u8BA1",
         "reference_code": "S0_ASSET", "balance_type": "Closing Balance",
         "calculation_formula": asset_f, "fieldtype": "Currency", "bold_text": 1},
        {"data_source": "Account Data", "display_name": "\u8D1F\u503A\u5408\u8BA1",
         "reference_code": "S0_LIAB", "balance_type": "Closing Balance",
         "calculation_formula": liab_f, "reverse_sign": 1, "fieldtype": "Currency"},
        {"data_source": "Column Break", "display_name": seg1_name},
        {"data_source": "Account Data", "display_name": "\u6240\u6709\u8005\u6743\u76CA\u5408\u8BA1",
         "reference_code": "S1_EQ", "balance_type": "Closing Balance",
         "calculation_formula": eq_f, "reverse_sign": 1, "fieldtype": "Currency",
         "bold_text": 1},
        {"data_source": "Account Data", "display_name": "\u8D1F\u503A\u5408\u8BA1(\u53F3\u680F)",
         "reference_code": "S1_LIAB", "balance_type": "Closing Balance",
         "calculation_formula": liab_f, "reverse_sign": 1, "fieldtype": "Currency"},
    ]


def baseline():
    return {
        "GL Entry": frappe.db.count("GL Entry"),
        "Stock Ledger Entry": frappe.db.count("Stock Ledger Entry"),
        "Account": frappe.db.count("Account"),
        "Company": frappe.db.count("Company"),
        "Financial Report Template": frappe.db.count("Financial Report Template"),
        "Fiscal Year": sorted(r.name for r in frappe.get_all("Fiscal Year", fields=["name"])),
        "Error Log": frappe.db.count("Error Log"),
        "Translation": frappe.db.count("Translation"),
        "Prepared Report": frappe.db.count("Prepared Report"),
        "Access Log": frappe.db.count("Access Log"),
        "File": frappe.db.count("File"),
    }


frappe.init(site=SITE)
frappe.connect()
frappe.set_user("Administrator")
os.makedirs(OUT_DIR, exist_ok=True)

_orig_enable_prepared_report = None
_orig_enqueue = None
report_mod = None
files_written = []

try:
    rec("baseline_before", baseline())

    from frappe.core.doctype.report import report as report_mod
    import frappe.desk.query_report as qr_mod
    from frappe.translate import get_all_translations, get_user_lang
    from erpnext.accounts.doctype.financial_report_template.financial_report_engine import (
        FinancialReportEngine,
    )

    # ---------- 0a. neutralize the prepared-report cross-connection commit ----------
    def _no_op_enable_prepared_report(report=None, site=None):
        rec("UNEXPECTED_enable_prepared_report_call", {"report": report, "site": site})

    _orig_enable_prepared_report = report_mod.enable_prepared_report
    report_mod.enable_prepared_report = _no_op_enable_prepared_report

    # ---------- 0b. tripwire on frappe.enqueue ----------
    # query_report.py:418 enqueues only under `if export_in_background:`. We never set
    # that, but assert it rather than assume it: any call lands in this list.
    enqueue_calls = []

    def _tripwire_enqueue(*a, **kw):
        enqueue_calls.append({"args": [str(x) for x in a], "kwargs": {k: str(v) for k, v in kw.items()}})
        raise AssertionError("frappe.enqueue was called on the export path -- probe aborted")

    _orig_enqueue = frappe.enqueue
    frappe.enqueue = _tripwire_enqueue
    rec("safety_patches_installed", {
        "enable_prepared_report": "no-op recorder",
        "frappe.enqueue": "tripwire that raises",
        "enqueue_site_on_path": [ln.strip() for ln in inspect.getsource(qr_mod).splitlines()
                                 if "export_in_background" in ln][:6],
    })

    # ---------- 1. lang: the render layer's answer depends on it ----------
    real_lang = get_user_lang("Administrator")
    frappe.local.lang = real_lang
    zh_dict = get_all_translations(real_lang)
    rec("lang_used_for_export", {
        "get_user_lang(Administrator)": real_lang,
        "frappe.local.lang_set_to": frappe.local.lang,
        "dict_entries": len(zh_dict),
        "note": "a real Desk export runs under the session user's lang; forced here to match",
    })

    # ---------- 2. POS control chosen from the LIVE dict ----------
    pos_candidates = ["Assets", "Liabilities", "Equity", "Total", "Account",
                      "Balance Sheet", "Opening", "Closing", "Currency", "Account Name"]
    pos_found = [{"key": k, "zh_value": zh_dict[k]}
                 for k in pos_candidates if k in zh_dict and zh_dict[k] != k]
    seg_pos_0 = pos_found[0]["key"] if pos_found else None
    seg_pos_1 = (pos_found[1]["key"] if len(pos_found) > 1 else
                 (pos_found[0]["key"] if pos_found else None))
    rec("pos_control_chosen", {
        "segment_0": seg_pos_0, "segment_0_zh": zh_dict.get(seg_pos_0),
        "segment_1": seg_pos_1, "segment_1_zh": zh_dict.get(seg_pos_1),
        "prediction": ("account column header == the bare key -> SUBSTITUTED in the file; "
                       "period column header == '<key> - 2026' -> not a key -> passes "
                       "through. Both inside one workbook."),
    })

    # ---------- 3. FY 2025 + templates, all inside the txn ----------
    if not frappe.db.exists("Fiscal Year", "2025"):
        fy = frappe.get_doc({"doctype": "Fiscal Year", "year": "2025",
                             "year_start_date": "2025-01-01", "year_end_date": "2025-12-31"})
        fy.insert(ignore_permissions=True)
        rec("created_fiscal_year", {"name": fy.name, "note": "inside txn, rolled back"})

    cells = [("NEG_statutory_chinese", TPL_NEG, SEG_NEG_0, SEG_NEG_1),
             ("LONG_realistic_panes", TPL_LONG, SEG_LONG_0, SEG_LONG_1)]
    if seg_pos_0:
        cells.append(("POS_translated_key", TPL_POS, seg_pos_0, seg_pos_1))

    built = {}
    for tag, name, s0, s1 in cells:
        d = frappe.get_doc({
            "doctype": "Financial Report Template", "template_name": name,
            "report_type": "Balance Sheet",   # module deliberately EMPTY
            "rows": bs_rows(s0, s1),
        })
        d.insert(ignore_permissions=True)
        built[tag] = d.name
        rec("template_built_" + tag, {"name": d.name, "module": d.module})

    def filters_for(tpl):
        return {
            "company": COMPANY, "report_template": tpl,
            "filter_based_on": "Fiscal Year",
            "from_fiscal_year": "2025", "to_fiscal_year": "2026",
            "periodicity": "Yearly", "selected_view": "Report",
            "include_default_book_entries": 1,
        }

    # ---------- 4. engine labels, as the comparison baseline ----------
    engine_labels = {}
    for tag, tpl in built.items():
        cols = FinancialReportEngine().execute(frappe._dict(filters_for(tpl)))[0]
        engine_labels[tag] = [{"fieldname": c.get("fieldname"), "label": c.get("label"),
                               "hidden": c.get("hidden", 0)} for c in cols]
        rec("ENGINE_LABELS_" + tag, engine_labels[tag])

    # ---------- 5. THE REAL EXPORT, then read the bytes back ----------
    import openpyxl

    from frappe.desk.reportview import clean_params, parse_json
    from frappe.desk.utils import pop_csv_params

    def prep(fp):
        """Reproduce export_query()'s own preprocessing (query_report.py:404-407):
        pop_csv_params -> clean_params -> parse_json. parse_json is what turns the
        `filters` JSON STRING into a dict; without it erpnext's get_xlsx_styles hook
        (financial_report_engine.py:1938 `metadata.filters.get(...)`) gets a str and
        raises. Skipping it was a defect in the first run of this probe, not a finding
        about the product."""
        csvp = pop_csv_params(fp)
        clean_params(fp)
        parse_json(fp)
        return csvp

    for tag, tpl in built.items():
        for include_hidden in (0, 1):
            form_params = frappe._dict({
                "report_name": REPORT_NAME,
                "file_format_type": "Excel",
                "filters": json.dumps(filters_for(tpl)),
                "applied_filters": json.dumps(filters_for(tpl)),
                "include_indentation": 0,
                "include_filters": 0,
                "include_hidden_columns": include_hidden,
                "custom_columns": "[]",
                "visible_idx": json.dumps([]),
                "ignore_visible_idx": 1,
            })
            csv_params = prep(form_params)
            key = "%s_hidden%d" % (tag, include_hidden)
            try:
                # populate_response=False -> bytes come back to us; nothing touches
                # frappe.response and no File row is created
                rname, ext, content = qr_mod._export_query(
                    form_params, csv_params=csv_params, populate_response=False)
            except Exception as e:
                rec("EXPORT_ERROR_" + key, {"err": str(e), "tb": traceback.format_exc()[-1500:]})
                continue

            path = os.path.join(OUT_DIR, "xlsx-%s.%s" % (key, ext))
            with open(path, "wb") as f:
                f.write(content)
            files_written.append(path)

            # independent reader -- openpyxl, not frappe's code
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_names = list(wb.sheetnames)
            ws = wb[sheet_names[0]]
            rows = []
            for r in ws.iter_rows(min_row=1, max_row=3, values_only=True):
                rows.append(list(r))
            wb.close()

            header = [c for c in (rows[0] if rows else [])]
            eng_visible = [c["label"] for c in engine_labels[tag]
                           if include_hidden or not c["hidden"]]
            rec("XLSX_HEADER_" + key, {
                "returned_report_name": rname,
                "file": path,
                "bytes": len(content),
                "sha256": hashlib.sha256(content).hexdigest()[:32],
                "sheet_names": sheet_names,
                "header_row_verbatim": header,
                "engine_labels_same_visibility": eng_visible,
                "byte_identical_to_engine": header == eng_visible,
                "first_data_rows": rows[1:3],
            })
            rec("XLSX_HEADER_CHARDUMP_" + key, [char_dump(h) for h in header])
            diffs = []
            for i, h in enumerate(header):
                exp = eng_visible[i] if i < len(eng_visible) else None
                if h != exp:
                    diffs.append({"index": i, "engine_label": exp, "in_xlsx": h,
                                  "engine_chars": char_dump(exp), "xlsx_chars": char_dump(h)})
            rec("XLSX_HEADER_DIFFS_" + key, {
                "count": len(diffs), "diffs": diffs,
                "discriminates": bool(diffs) or None,
            })

    # ---------- 5b. CSV export: the same build_xlsx_data header row ----------
    # Same _() transform (build_xlsx_data is shared); included because CSV is the other
    # format the Export dialog offers and it is cheap to settle in the same run.
    for tag, tpl in built.items():
        form_params = frappe._dict({
            "report_name": REPORT_NAME, "file_format_type": "CSV",
            "filters": json.dumps(filters_for(tpl)),
            "applied_filters": json.dumps(filters_for(tpl)),
            "include_indentation": 0, "include_filters": 0,
            "include_hidden_columns": 0, "custom_columns": "[]",
            "visible_idx": json.dumps([]), "ignore_visible_idx": 1,
        })
        csv_params = prep(form_params)
        try:
            rname, ext, content = qr_mod._export_query(
                form_params, csv_params=csv_params, populate_response=False)
        except Exception as e:
            rec("CSV_EXPORT_ERROR_" + tag, {"err": str(e)})
            continue
        path = os.path.join(OUT_DIR, "csv-%s.%s" % (tag, ext))
        with open(path, "wb") as f:
            f.write(content)
        files_written.append(path)
        text = content.decode("utf-8-sig")
        first_line = text.splitlines()[0] if text.splitlines() else ""
        eng_visible = [c["label"] for c in engine_labels[tag] if not c["hidden"]]
        rec("CSV_HEADER_" + tag, {
            "file": path, "bytes": len(content),
            "header_line_verbatim": first_line,
            "engine_labels_visible": eng_visible,
            "all_engine_labels_present_verbatim": all(lbl in first_line for lbl in eng_visible),
            "encoding_note": "decoded as utf-8-sig; BOM presence: %s" % content[:3].hex(),
        })

    # sheet name (report title) also passes through _() in provide_binary_file
    rec("sheet_name_sanitizer", {
        "MAX_SHEET_NAME_LENGTH": 31,
        "INVALID_SHEET_NAME_RE": r"[\[\]:*?/\\]",
        "note": ("xlsxutils.get_sanitized_sheet_name truncates the SHEET NAME to 31 chars "
                 "and strips []:*?/\\ -- it touches the tab name, not the header cells"),
    })
    rec("enqueue_tripwire_calls", enqueue_calls)

except Exception as e:
    rec("probe_error", {"err": str(e), "tb": traceback.format_exc()[-3000:]})

finally:
    if _orig_enqueue is not None:
        frappe.enqueue = _orig_enqueue
    if _orig_enable_prepared_report is not None and report_mod is not None:
        report_mod.enable_prepared_report = _orig_enable_prepared_report

    frappe.db.rollback()

    rec("baseline_after_rollback", baseline())
    rec("templates_exist_after_rollback", {
        n: bool(frappe.db.exists("Financial Report Template", n))
        for n in [TPL_NEG, TPL_LONG, TPL_POS]})
    rec("Report_prepared_report_after_rollback",
        frappe.db.get_value("Report", REPORT_NAME, "prepared_report"))
    # reported, never deleted -- MyISAM
    rec("error_log_rows_now", frappe.get_all(
        "Error Log", fields=["name", "method", "creation"], limit=20))
    rec("access_log_rows_now", frappe.db.count("Access Log"))

    rec("files_this_probe_wrote", files_written)
    # prove nothing landed in frappe's private files dir
    priv = frappe.get_site_path("private", "files")
    pub = frappe.get_site_path("public", "files")
    strays = []
    for d in (priv, pub):
        for pat in ("*V31*", "*v31*", "*PROBE*", "*probe*", "*ZZ*", "*Custom Financial*"):
            strays += glob.glob(os.path.join(d, pat))
    rec("stray_files_in_site_files_dirs", sorted(set(strays)))
    tpl_strays = []
    for pat in ("*probe*", "*PROBE*", "*v31*", "*V31*", "*zz*", "*ZZ*"):
        tpl_strays += glob.glob(frappe.get_app_path(
            "erpnext", "accounts", "financial_report_template", pat))
    rec("stray_exported_templates", sorted(set(tpl_strays)))

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2, default=str)
    print("\nWROTE " + OUT, flush=True)
